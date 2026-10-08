import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from tools import studio


class StudioTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        shutil.copytree(studio.ROOT / 'templates', self.root / 'templates')

    def test_templates_and_no_overwrite(self):
        with patch.object(studio, 'ROOT', self.root):
            for kind in studio.KINDS:
                studio.new_project(kind, kind)
                brief = self.root / 'projects' / kind / 'brief.md'
                brief.write_text('已修改的原稿', encoding='utf-8')
                with self.assertRaises(FileExistsError):
                    studio.new_project(kind, kind)
                self.assertEqual(brief.read_text(encoding='utf-8'), '已修改的原稿')
            for name in ('..', '../逃逸', 'a/b', 'a\\b', ' '):
                with self.assertRaises(ValueError):
                    studio.new_project('tvc', name)

    def test_storyboard_validation(self):
        data = json.loads((self.root / 'templates/storyboard.json').read_text())
        path = self.root / 'shots.json'
        path.write_text(json.dumps(data))
        studio.load_storyboard(path)
        for duration in (0, -1, True, float('inf'), float('nan')):
            data['shots'][0]['duration'] = duration
            path.write_text(json.dumps(data))
            with self.assertRaises(ValueError):
                studio.load_storyboard(path)
        data['shots'][0]['duration'] = 3
        data['shots'].append(dict(data['shots'][0]))
        path.write_text(json.dumps(data))
        with self.assertRaises(ValueError):
            studio.load_storyboard(path)

    def test_cli_refuses_existing_output(self):
        dest = self.root / '稿件.docx'
        dest.write_bytes(b'original')
        result = subprocess.run([sys.executable, str(studio.ROOT / 'tools/studio.py'), 'docx',
                                 str(self.root / 'templates/tvc.md'), str(dest)], capture_output=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(dest.read_bytes(), b'original')

    def test_office_exports(self):
        from docx import Document
        from pptx import Presentation
        docx = self.root / '样例.docx'
        pptx = self.root / '样例.pptx'
        studio.export_docx(self.root / 'templates/tvc.md', docx)
        studio.export_pptx(self.root / 'templates/storyboard.json', pptx)
        self.assertTrue(any('TVC' in p.text for p in Document(docx).paragraphs))
        slides = Presentation(pptx).slides
        data = studio.load_storyboard(self.root / 'templates/storyboard.json')
        self.assertEqual(len(slides), len(data['shots']))
        self.assertTrue(any('画面待补' in s.text for s in slides[0].shapes if s.has_text_frame))

    def test_storyboard_image_keeps_aspect_ratio(self):
        from PIL import Image
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        Image.new('RGB', (600, 900), 'white').save(self.root / 'portrait.png')
        data = studio.load_storyboard(self.root / 'templates/storyboard.json')
        data['shots'][0]['image'] = 'portrait.png'
        source = self.root / 'shots.json'
        source.write_text(json.dumps(data), encoding='utf-8')
        dest = self.root / 'portrait.pptx'
        studio.export_pptx(source, dest)
        pictures = [s for s in Presentation(dest).slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(pictures), 1)
        self.assertAlmostEqual(pictures[0].width / pictures[0].height, 2 / 3, places=5)
        self.assertEqual(pictures[0].crop_left + pictures[0].crop_right, 0)


if __name__ == '__main__':
    unittest.main()
