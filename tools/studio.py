#!/usr/bin/env python3
"""项目模板、统计、文本比较和基础 Office 导出。"""
import argparse
import difflib
import importlib.metadata
import json
from pathlib import Path
import re
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
KINDS = ('scifi', 'tvc', 'corporate', 'documentary', 'screenplay')

def doctor(require_render=False):
    ok = True
    for pkg in ('python-docx', 'python-pptx', 'Pillow', 'pypdf'):
        try:
            print(f'{pkg}: {importlib.metadata.version(pkg)}')
        except importlib.metadata.PackageNotFoundError:
            print(f'{pkg}: 未安装'); ok = False
    for tool in ('libreoffice', 'soffice', 'pdftoppm', 'fc-match'):
        print(f'{tool}: {shutil.which(tool) or "未发现"}')
    if require_render:
        ok = ok and bool(shutil.which('libreoffice') or shutil.which('soffice')) and bool(shutil.which('pdftoppm'))
    return 0 if ok else 1

def new_project(kind, name):
    if name in ('.', '..') or not name.strip() or any(c in name for c in '/\\'):
        raise ValueError('项目名不能为空或包含路径分隔符')
    target = ROOT / 'projects' / name
    target.mkdir(parents=True, exist_ok=False)
    shutil.copyfile(ROOT / 'templates' / f'{kind}.md', target / 'brief.md')
    for folder in ('drafts', 'revisions', 'research', 'storyboards'):
        (target / folder).mkdir()
    shutil.copyfile(ROOT / 'templates/revision.md', target / 'revisions/修改记录.md')
    shutil.copyfile(ROOT / 'templates/sources.csv', target / 'research/来源.csv')
    print(target)

def load_storyboard(path):
    data = json.loads(path.read_text(encoding='utf-8'))
    if not isinstance(data.get('shots'), list) or not data['shots']:
        raise ValueError('shots 必须是非空列表')
    ids = set()
    for shot in data['shots']:
        sid = str(shot.get('id', '')).strip()
        if not sid or sid in ids:
            raise ValueError('镜号必须存在且唯一')
        ids.add(sid)
        duration = shot.get('duration')
        if isinstance(duration, bool) or not isinstance(duration, (int,float)) or not 0 < duration < float('inf'):
            raise ValueError(f'镜头 {sid} 的时长必须是有限正数')
        for field in ('visual', 'dialogue', 'sound'):
            text = shot.get(field, '')
            if not isinstance(text, str) or len(text) > 100 or text.count('\n') > 3:
                raise ValueError(f'镜头 {sid} 的 {field} 过长或格式错误，请拆镜或精简')
        for field in ('id', 'scene', 'size', 'camera', 'status'):
            if len(str(shot.get(field, ''))) > 24:
                raise ValueError(f'镜头 {sid} 的 {field} 过长')
        if shot.get('image') and not (path.parent / shot['image']).is_file():
            raise ValueError(f'镜头 {sid} 的图片不存在')
    return data

def export_pptx(source, dest):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from PIL import Image
    data = load_storyboard(source)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    def text(slide, x, y, w, h, value, size):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame; tf.word_wrap = True
        for n, line in enumerate(str(value).splitlines() or ['']):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = 'Noto Sans CJK SC'; p.font.size = Pt(size)
            p.font.color.rgb = RGBColor.from_string('18252D')
        return box
    for shot in data['shots']:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        text(s,.55,.25,12.2,.65,f"镜头 {shot['id']} · 场 {shot.get('scene','')} · {shot['duration']} 秒",32)
        text(s,.55,1.05,7.5,.45,f"{shot.get('size','')} / {shot.get('camera','')}",18)
        if shot.get('image'):
            image = source.parent / shot['image']
            with Image.open(image) as im:
                iw, ih = im.size
            scale = min(7.5/iw, 4.22/ih)
            w,h = iw*scale, ih*scale
            s.shapes.add_picture(str(image), Inches(.55+(7.5-w)/2), Inches(1.65+(4.22-h)/2), width=Inches(w), height=Inches(h))
        else:
            text(s,1.8,3.0,5,1,'画面待补',28)
        for y, label, field in ((1.65,'画面','visual'),(3.30,'台词／旁白','dialogue'),(4.95,'声音','sound')):
            text(s,8.4,y,4.25,.4,label,18)
            text(s,8.4,y+.42,4.25,1.18,shot.get(field,''),18)
        text(s,.55,6.75,12.2,.4,f"素材状态：{shot.get('status','待确认')} · 总时长 {sum(x['duration'] for x in data['shots']):g} 秒",17)
    prs.save(dest)

def export_docx(source, dest):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for style in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        st = doc.styles[style]
        st.font.name = 'Noto Sans CJK SC'
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    doc.styles['Normal'].font.size = Pt(11)
    for line in source.read_text(encoding='utf-8').splitlines():
        if not line.strip(): continue
        if re.match(r'^#{1,4} ',line):
            count = len(line)-len(line.lstrip('#'))
            doc.add_heading(line[count:].strip(), level=0 if count == 1 else min(count-1,3))
        else:
            doc.add_paragraph(line)
    doc.save(dest)

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='cmd', required=True)
    p = sub.add_parser('new'); p.add_argument('kind',choices=KINDS); p.add_argument('name')
    p = sub.add_parser('doctor'); p.add_argument('--require-render',action='store_true')
    p = sub.add_parser('stats'); p.add_argument('path',type=Path)
    p = sub.add_parser('diff'); p.add_argument('old',type=Path); p.add_argument('new',type=Path)
    for command in ('docx','pptx'):
        p = sub.add_parser(command); p.add_argument('source',type=Path); p.add_argument('dest',type=Path)
    args = parser.parse_args()
    if args.cmd == 'doctor': return doctor(args.require_render)
    if args.cmd == 'new': new_project(args.kind,args.name)
    if args.cmd == 'stats':
        text = args.path.read_text(encoding='utf-8')
        print(json.dumps({'汉字':len(re.findall(r'[\u3400-\u9fff]',text)), '英文单词':len(re.findall(r'[A-Za-z]+',text)), '非空白字符':len(re.sub(r'\s','',text))},ensure_ascii=False))
    if args.cmd == 'diff':
        print(''.join(difflib.unified_diff(args.old.read_text(encoding='utf-8').splitlines(True),args.new.read_text(encoding='utf-8').splitlines(True),fromfile=str(args.old),tofile=str(args.new))),end='')
    if args.cmd in ('docx','pptx'):
        if args.dest.exists(): raise ValueError('目标已存在，请使用新版本文件名')
        args.dest.parent.mkdir(parents=True, exist_ok=True)
        (export_docx if args.cmd == 'docx' else export_pptx)(args.source,args.dest)
        print(args.dest)
    return 0

if __name__ == '__main__':
    try:
        sys.exit(main())
    except (ValueError,OSError,KeyError,TypeError) as exc:
        print(f'错误：{exc}',file=sys.stderr); sys.exit(1)
